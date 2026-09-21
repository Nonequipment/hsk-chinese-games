"""Extract the checked HSK workbook into deterministic application JSON."""
from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path
import openpyxl


def iso(value: date | datetime | None) -> str | None:
    return value.date().isoformat() if isinstance(value, datetime) else value.isoformat() if value else None


def extract(source: Path) -> dict:
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    vocab_rows = list(workbook["Vocabulary_1200"].iter_rows(min_row=2, values_only=True))
    set_rows = list(workbook["Sets_60"].iter_rows(min_row=2, values_only=True))
    day_rows = list(workbook["Plan_26_Days"].iter_rows(min_row=2, values_only=True))

    items = [{
        "id": f"V{int(row[0]):04d}", "number": int(row[0]), "setId": str(row[1]),
        "category": str(row[2]).strip(), "scheduledDate": iso(row[3]),
        "hanzi": str(row[4]).strip(), "pinyin": str(row[5]).strip(), "thai": str(row[6]).strip(),
    } for row in vocab_rows if row[0] is not None]

    by_set: dict[str, list[str]] = {}
    for item in items:
        by_set.setdefault(item["setId"], []).append(item["id"])

    sets = [{
        "id": str(row[0]), "category": str(row[1]).strip(), "scheduledDate": iso(row[3]),
        "itemIds": by_set.get(str(row[0]), []),
    } for row in set_rows if row[0]]
    days = [{
        "day": int(row[0]), "date": iso(row[1]),
        "setIds": [value.strip() for value in str(row[2]).split(",")],
        "setCount": int(row[3]), "wordCount": int(row[4]),
    } for row in day_rows if row[0] is not None]

    assert len(items) == 1200, f"expected 1200 items, got {len(items)}"
    assert len({item['id'] for item in items}) == 1200, "vocabulary IDs must be unique"
    assert len(sets) == 60, f"expected 60 sets, got {len(sets)}"
    assert all(len(value) == 20 for value in by_set.values()), "every set must contain 20 items"
    assert all(item["hanzi"] and item["pinyin"] and item["thai"] for item in items), "empty required vocabulary value"
    assert days and days[0]["setIds"] == ["S01", "S02", "S03"], "Day 1 must contain S01-S03"
    return {"version": "2026-09-21", "targetDate": "2026-10-17", "items": items, "sets": sets, "days": days}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    data = extract(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"wrote {len(data['items'])} words in {len(data['sets'])} sets to {args.output}")


if __name__ == "__main__":
    main()
